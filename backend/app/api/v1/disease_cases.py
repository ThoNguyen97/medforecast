"""API số ca bệnh dịch tễ: tra cứu, nhập tay, import CSV, xuất Excel và ước tính nhu cầu thuốc theo ca."""
import csv
import io
import logging
from typing import List, Optional, Any
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException, status, Request, Query, UploadFile, File
from sqlalchemy.orm import Session
from sqlalchemy import func

from app.database import get_db
from app.dependencies import get_current_user
from app.models.user import User
from app.schemas.base import DiseaseCaseCreate, DiseaseCaseResponse, DiseaseCaseUpdate
from app.services.disease_case_service import DiseaseCaseService

from app.utils.icd_groups import NHOM_ICD, dieu_kien_nhom, nhom_cua_ma

router = APIRouter()
logger = logging.getLogger(__name__)


def get_client_ip(request: Request) -> str:
    """Lấy IP client, ưu tiên X-Forwarded-For khi chạy sau proxy."""
    forwarded = request.headers.get("X-Forwarded-For")
    if forwarded:
        return forwarded.split(",")[0].strip()
    return request.client.host if request.client else "unknown"


@router.get("/", response_model=List[DiseaseCaseResponse])
def list_disease_cases(
    skip: int = Query(0, ge=0),
    limit: int = Query(100, ge=1),  # không đặt trần: giao diện có thể lấy toàn bộ
    disease_type: Optional[str] = Query(None, description="Lọc theo mã ICD 3 ký tự (J00…J22)"),
    disease_group: Optional[str] = Query(None, description="Lọc theo NHÓM ICD: J00-J06 / J09-J18 / J20-J22"),
    location: Optional[str] = Query(None, description="Filter by location (Tỉnh/Thành)"),
    start_date: Optional[str] = Query(None, description="Start date filter (YYYY-MM-DD format)"),
    end_date: Optional[str] = Query(None, description="End date filter (YYYY-MM-DD format)"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
) -> Any:
    """Danh sách bản ghi số ca, lọc theo nhóm ICD, mã ICD, tỉnh/thành và khoảng ngày."""
    from app.models.disease_case import DiseaseCase

    q = db.query(DiseaseCase)
    if disease_group:
        # OR thêm dải mã để không bỏ sót bản ghi chưa điền cột disease_group.
        q = q.filter(dieu_kien_nhom(DiseaseCase, disease_group))
    if disease_type:
        q = q.filter(DiseaseCase.icd_code == disease_type)
    if location:
        q = q.filter(DiseaseCase.location == location)
    if start_date:
        try:
            start_dt = datetime.fromisoformat(start_date)
            q = q.filter(DiseaseCase.recorded_at >= start_dt)
        except ValueError:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Invalid start_date format: {start_date}. Expected YYYY-MM-DD."
            )
    if end_date:
        try:
            end_dt = datetime.fromisoformat(end_date)
            # Mốc cuối bao trọn ngày end_date (đến 23:59:59).
            from datetime import timedelta
            end_dt = end_dt + timedelta(days=1) - timedelta(seconds=1)
            q = q.filter(DiseaseCase.recorded_at <= end_dt)
        except ValueError:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Invalid end_date format: {end_date}. Expected YYYY-MM-DD."
            )
    return (
        q.order_by(DiseaseCase.recorded_at.desc())
        .offset(skip)
        .limit(limit)
        .all()
    )


# ── Xuất Excel ──
def _khoang_thang(start_month: Optional[str], end_month: Optional[str]):
    """'YYYY-MM' → (đầu tháng đầu, cuối tháng cuối); đảo nếu nhập ngược, (None, None) khi không lọc."""
    from datetime import timedelta

    def _parse(s: Optional[str]) -> Optional[datetime]:
        if not s:
            return None
        try:
            return datetime.strptime(s.strip(), "%Y-%m")
        except ValueError:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Tháng không hợp lệ: {s}. Định dạng YYYY-MM.",
            )

    tu, den = _parse(start_month), _parse(end_month)
    if tu and den and tu > den:
        tu, den = den, tu
    if den:
        nam, thang = (den.year + 1, 1) if den.month == 12 else (den.year, den.month + 1)
        den = datetime(nam, thang, 1) - timedelta(seconds=1)
    return tu, den


@router.get("/export")
def export_disease_cases_excel(
    disease_group: Optional[str] = Query(None, description="Nhóm ICD: J00-J06 / J09-J18 / J20-J22"),
    location: Optional[str] = Query(None, description="Tỉnh/Thành (chấp nhận mọi biến thể tên)"),
    start_month: Optional[str] = Query(None, description="Từ tháng, YYYY-MM"),
    end_month: Optional[str] = Query(None, description="Đến tháng, YYYY-MM"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> Any:
    """Xuất số ca bệnh ra .xlsx (Chi tiết, Tổng hợp theo tháng × nhóm × tỉnh, Bộ lọc) theo bộ lọc giao diện."""
    import io as _io
    from fastapi.responses import Response
    from app.models.disease_case import DiseaseCase
    from app.utils.province_alias import province_aliases, ten_chuan

    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(
            status_code=500, detail="openpyxl chưa được cài — không thể xuất Excel",
        )

    q = db.query(DiseaseCase)
    if disease_group:
        q = q.filter(dieu_kien_nhom(DiseaseCase, disease_group))
    if location:
        q = q.filter(DiseaseCase.location.in_(province_aliases(location)))
    tu, den = _khoang_thang(start_month, end_month)
    if tu:
        q = q.filter(DiseaseCase.recorded_at >= tu)
    if den:
        q = q.filter(DiseaseCase.recorded_at <= den)

    rows = q.order_by(
        DiseaseCase.recorded_at.desc(), DiseaseCase.location, DiseaseCase.icd_code
    ).all()

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    bold = Font(bold=True)

    def _ghi_header(ws, headers, widths):
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=i, value=h)
            c.fill, c.font, c.alignment = header_fill, header_font, center
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    wb = openpyxl.Workbook()

    # ── Sheet 1: Chi tiết ──
    ws = wb.active
    ws.title = "Chi tiết"
    _ghi_header(
        ws,
        ["STT", "Tháng/Năm", "Mã nhóm", "Nhóm bệnh", "Mã ICD", "Tên bệnh",
         "Tỉnh/Thành", "Số ca mắc", "Nguồn dữ liệu", "Ghi chú"],
        [6, 11, 10, 38, 9, 42, 22, 11, 16, 30],
    )

    tong_hop: dict = {}  # (yyyy-mm, nhóm, tỉnh) -> tổng ca
    tong_ca = 0
    for i, r in enumerate(rows, 1):
        nhom = r.disease_group or nhom_cua_ma(r.icd_code or "") or ""
        thang = r.recorded_at.strftime("%m/%Y") if r.recorded_at else ""
        ky = r.recorded_at.strftime("%Y-%m") if r.recorded_at else ""
        tinh = ten_chuan(r.location or "")
        so_ca = int(r.case_count or 0)
        tong_ca += so_ca

        ws.append([
            i, thang, nhom, NHOM_ICD.get(nhom, nhom or "—"), r.icd_code,
            r.disease_name, tinh, so_ca, r.data_source or "", r.note or "",
        ])
        k = (ky, nhom, tinh)
        tong_hop[k] = tong_hop.get(k, 0) + so_ca

    r_tong = ws.max_row + 1
    ws.cell(row=r_tong, column=7, value="TỔNG").font = bold
    c = ws.cell(row=r_tong, column=8, value=tong_ca)
    c.font = bold
    for r_idx in range(2, r_tong + 1):
        ws.cell(row=r_idx, column=1).alignment = center
        ws.cell(row=r_idx, column=2).alignment = center
        ws.cell(row=r_idx, column=8).number_format = "#,##0"

    # ── Sheet 2: Tổng hợp ──
    ws2 = wb.create_sheet("Tổng hợp")
    _ghi_header(
        ws2,
        ["Tháng/Năm", "Mã nhóm", "Nhóm bệnh", "Tỉnh/Thành", "Số ca mắc"],
        [11, 10, 38, 22, 11],
    )
    for (ky, nhom, tinh), so_ca in sorted(
        tong_hop.items(), key=lambda kv: (kv[0][0], kv[0][1], kv[0][2]), reverse=False
    ):
        thang = f"{ky[5:7]}/{ky[0:4]}" if len(ky) == 7 else ky
        ws2.append([thang, nhom, NHOM_ICD.get(nhom, nhom or "—"), tinh, so_ca])
    r2 = ws2.max_row + 1
    ws2.cell(row=r2, column=4, value="TỔNG").font = bold
    ws2.cell(row=r2, column=5, value=tong_ca).font = bold
    for r_idx in range(2, r2 + 1):
        ws2.cell(row=r_idx, column=1).alignment = center
        ws2.cell(row=r_idx, column=5).number_format = "#,##0"

    # ── Sheet 3: Bộ lọc đã áp dụng ──
    ws3 = wb.create_sheet("Bộ lọc")
    ws3.column_dimensions["A"].width = 22
    ws3.column_dimensions["B"].width = 40
    for label, val in (
        ("Thời điểm xuất", datetime.now().strftime("%d/%m/%Y %H:%M:%S")),
        ("Người xuất", getattr(current_user, "username", "") or ""),
        ("Nhóm bệnh", NHOM_ICD.get(disease_group or "", disease_group) or "Tất cả"),
        ("Tỉnh/Thành", ten_chuan(location) if location else "Tất cả"),
        ("Từ tháng", f"{start_month[5:7]}/{start_month[0:4]}" if start_month else "—"),
        ("Đến tháng", f"{end_month[5:7]}/{end_month[0:4]}" if end_month else "—"),
        ("Số bản ghi", len(rows)),
        ("Tổng số ca", tong_ca),
    ):
        ws3.append([label, val])
        ws3.cell(row=ws3.max_row, column=1).font = bold

    buf = _io.BytesIO()
    wb.save(buf)
    filename = f"du_lieu_benh_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/", response_model=DiseaseCaseResponse, status_code=status.HTTP_201_CREATED)
def create_disease_case(
    data: DiseaseCaseCreate,
    request: Request,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
) -> Any:
    """Thêm một bản ghi số ca; trùng (tháng, mã ICD, khu vực) thì 409 kèm `existing_id`."""
    from app.models.disease_case import DiseaseCase
    from sqlalchemy import extract

    service = DiseaseCaseService(db)
    client_ip = get_client_ip(request)

    # Không so disease_type: mọi bản ghi đều là "respiratory".
    existing = (
        db.query(DiseaseCase)
        .filter(
            extract("year", DiseaseCase.recorded_at) == data.recorded_at.year,
            extract("month", DiseaseCase.recorded_at) == data.recorded_at.month,
            DiseaseCase.icd_code == data.icd_code,
            DiseaseCase.location == data.location,
        )
        .first()
    )
    if existing:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail={
                "message": "Đã tồn tại bản ghi cho cùng tháng, bệnh và khu vực. "
                           "Vui lòng cập nhật thay vì thêm mới.",
                "existing_id": existing.id,
            },
        )

    try:
        new_case = service.create_disease_case(
            data=data,
            created_by_user_id=current_user.id,
            ip_address=client_ip
        )
        new_case.note = data.note
        new_case.created_by = current_user.username
        db.commit()
        db.refresh(new_case)
        return new_case
    except Exception as e:
        logger.error(f"Failed to create disease case: {str(e)}")
        raise


@router.get("/stats")
def get_disease_case_statistics(
    location: Optional[str] = Query(None, description="Filter by location"),
    start_date: Optional[datetime] = Query(None, description="Start date (ISO format)"),
    end_date: Optional[datetime] = Query(None, description="End date (ISO format)"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
) -> Any:
    """Thống kê theo loại bệnh: tổng số ca, số bản ghi và ngày ghi nhận gần nhất."""
    service = DiseaseCaseService(db)
    statistics = service.get_statistics(
        location=location,
        start_date=start_date,
        end_date=end_date
    )
    return {
        "statistics": statistics,
        "filters": {
            "location": location,
            "start_date": start_date.isoformat() if start_date else None,
            "end_date": end_date.isoformat() if end_date else None
        }
    }


@router.get("/trends")
def get_disease_case_trends(
    disease_type: Optional[str] = Query(None, description="Filter by disease type"),
    location: Optional[str] = Query(None, description="Filter by location"),
    start_date: Optional[datetime] = Query(None, description="Start date (ISO format)"),
    end_date: Optional[datetime] = Query(None, description="End date (ISO format)"),
    limit: int = Query(100, ge=1, le=1000, description="Maximum number of data points"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
) -> Any:
    """Chuỗi thời gian số ca cộng dồn theo ngày."""
    service = DiseaseCaseService(db)
    trends = service.get_trends(
        disease_type=disease_type,
        location=location,
        start_date=start_date,
        end_date=end_date,
        limit=limit
    )
    return {
        "trends": trends,
        "filters": {
            "disease_type": disease_type,
            "location": location,
            "start_date": start_date.isoformat() if start_date else None,
            "end_date": end_date.isoformat() if end_date else None
        }
    }


# Tên bệnh tiếng Việt (hoặc mã ICD) → mã ICD, dùng cho import CSV; chỉ nhận 4 mã.
NHOM_BENH_MAPPING = {
    "Viêm phế quản cấp": "J20",
    "Nhiễm trùng đường hô hấp trên cấp": "J06",
    "Nhiễm trùng hô hấp trên cấp": "J06",
    "Viêm họng cấp": "J02",
    "Viêm xoang cấp": "J01",
    "J20": "J20",
    "J06": "J06",
    "J02": "J02",
    "J01": "J01",
}

# Mã ICD → tên bệnh chuẩn.
ICD_TO_NAME = {
    "J20": "Viêm phế quản cấp",
    "J06": "Nhiễm trùng đường hô hấp trên cấp",
    "J02": "Viêm họng cấp",
    "J01": "Viêm xoang cấp",
}


@router.post("/import-csv", status_code=status.HTTP_200_OK)
async def import_disease_cases_csv(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> Any:
    """Import số ca bệnh từ CSV (định dạng file mẫu hoặc CSV bệnh viện); khu vực mới tự đăng ký vào admin.regions."""
    from app.models.disease_case import DiseaseCase

    if not file.filename or not file.filename.lower().endswith((".csv",)):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File must be a CSV",
        )

    try:
        raw = await file.read()
        text = raw.decode("utf-8-sig", errors="replace")
        reader = list(csv.DictReader(io.StringIO(text)))
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Cannot parse CSV: {exc}",
        )

    if not reader:
        return {"status": "ok", "imported": 0, "updated": 0, "skipped": 0, "message": "Empty CSV"}

    headers = {h.strip() for h in (reader[0].keys() if reader else [])}

    imported = 0
    updated = 0
    skipped = 0
    errors: list[dict] = []

    new_regions_to_register: set[str] = set()

    def _register_pending_regions() -> None:
        """Thêm khu vực mới vào admin.regions (idempotent)."""
        if not new_regions_to_register:
            return
        from app.models.system_config import SystemConfig
        import json as _json

        cfg = (
            db.query(SystemConfig)
            .filter(SystemConfig.config_key == "admin.regions")
            .first()
        )
        if cfg is None:
            cfg = SystemConfig(
                config_key="admin.regions",
                config_value="[]",
                description="Danh mục khu vực",
            )
            db.add(cfg)
            db.flush()

        try:
            existing = _json.loads(cfg.config_value or "[]")
            if not isinstance(existing, list):
                existing = []
        except Exception:
            existing = []

        existing_names = {r.get("name") for r in existing if isinstance(r, dict)}
        added_count = 0
        for name in new_regions_to_register:
            if name in existing_names:
                continue
            existing.append({"name": name, "province": None, "description": ""})
            added_count += 1

        if added_count:
            cfg.config_value = _json.dumps(existing, ensure_ascii=False)
            cfg.updated_by = current_user.id
            logger.info(
                "CSV import: auto-registered %d new region(s): %s",
                added_count,
                ", ".join(sorted(new_regions_to_register)),
            )

    has_simple_format = (
        "month" in headers and 
        "region" in headers and 
        "cases" in headers and 
        ("disease_name" in headers or "disease_code" in headers)
    )
    
    if has_simple_format:
        # Định dạng file mẫu; cột supply_* chỉ để gộp dòng, không lưu (nhu cầu thuốc lấy từ Tầng 2).
        has_supply_col = (
            "supply_name" in headers
            or "supply_code" in headers
            or "drug_code" in headers
        )

        grouped: dict[tuple, dict] = {}
        for idx, row in enumerate(reader, start=2):
            month_str = (row.get("month") or "").strip()
            disease_code = (row.get("disease_code") or "").strip().upper()
            disease_name = (row.get("disease_name") or "").strip()
            disease = disease_code if disease_code else disease_name
            
            region = (row.get("region") or "").strip()
            cases_raw = (row.get("cases") or "").strip()
            note = (row.get("note") or "").strip() or None
            supply_name = (row.get("supply_name") or "").strip() if has_supply_col else ""
            supply_code = (row.get("supply_code") or "").strip() if has_supply_col else ""
            drug_code = (row.get("drug_code") or "").strip() if has_supply_col else ""
            supply_qty_raw = (row.get("supply_quantity") or "").strip() if has_supply_col else ""
            supply_unit = (row.get("supply_unit") or "").strip() if has_supply_col else ""
            supply_category = (row.get("supply_category") or "").strip() if has_supply_col else ""

            if not month_str or not disease or not region or not cases_raw:
                skipped += 1
                errors.append({
                    "row": idx,
                    "reason": "Thiếu cột bắt buộc (month / disease_code hoặc disease_name / region / cases)",
                    "data": {"month": month_str, "disease": disease, "region": region, "cases": cases_raw},
                })
                continue

            parsed: Optional[datetime] = None
            for fmt in ("%m/%Y", "%Y-%m", "%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
                try:
                    parsed = datetime.strptime(month_str, fmt)
                    break
                except ValueError:
                    continue
            if parsed is None:
                skipped += 1
                errors.append({
                    "row": idx,
                    "reason": f"Không hiểu định dạng tháng: '{month_str}' (chấp nhận MM/YYYY, YYYY-MM)",
                })
                continue
            recorded_at = datetime(parsed.year, parsed.month, 1)

            try:
                cases_int = int(float(cases_raw))
            except ValueError:
                skipped += 1
                errors.append({"row": idx, "reason": f"Số ca không hợp lệ: '{cases_raw}'"})
                continue
            if cases_int < 0:
                skipped += 1
                errors.append({"row": idx, "reason": f"Số ca phải >= 0, nhận được {cases_int}"})
                continue

            if disease in ICD_TO_NAME:
                disease_key = disease
            else:
                disease_key = NHOM_BENH_MAPPING.get(disease, disease)
            
            if disease_key not in ICD_TO_NAME:
                skipped += 1
                errors.append({
                    "row": idx,
                    "reason": f"Bệnh không hợp lệ: '{disease}'. Chỉ hỗ trợ 4 bệnh: {', '.join(ICD_TO_NAME.values())} hoặc mã ICD: {', '.join(ICD_TO_NAME.keys())}",
                })
                continue
            disease_name_std = ICD_TO_NAME[disease_key]
            key = (recorded_at, disease_key, region)

            entry = grouped.get(key)
            if entry is None:
                entry = {
                    "recorded_at": recorded_at,
                    "disease_key": disease_key,
                    "disease_name": disease_name_std,
                    "region": region,
                    "case_count": cases_int,
                    "note": note,
                    "supplies": [],
                    "first_row": idx,
                }
                grouped[key] = entry
            else:
                # Cùng bản ghi lặp ở nhiều dòng: giữ số ca lớn nhất.
                if cases_int > entry["case_count"]:
                    entry["case_count"] = cases_int
                if note and not entry["note"]:
                    entry["note"] = note

            if supply_name or supply_code or drug_code:
                try:
                    supply_qty_int = int(float(supply_qty_raw)) if supply_qty_raw else 0
                except ValueError:
                    skipped += 1
                    errors.append({
                        "row": idx,
                        "reason": f"Số lượng thuốc không hợp lệ: '{supply_qty_raw}'",
                    })
                    continue
                if supply_qty_int < 0:
                    skipped += 1
                    errors.append({
                        "row": idx,
                        "reason": f"Số lượng thuốc phải >= 0, nhận được {supply_qty_int}",
                    })
                    continue
                entry["supplies"].append({
                    "name": supply_name,
                    "supply_code": supply_code,
                    "drug_code": drug_code,
                    "qty": supply_qty_int,
                    "unit": supply_unit or "đơn vị",
                    "category": supply_category or "Khác",
                })

        for entry in grouped.values():
            recorded_at = entry["recorded_at"]
            disease_key = entry["disease_key"]
            disease_name_std = entry["disease_name"]
            region = entry["region"]
            cases_int = entry["case_count"]
            note = entry["note"]

            existing = (
                db.query(DiseaseCase)
                .filter(
                    DiseaseCase.recorded_at == recorded_at,
                    DiseaseCase.icd_code == disease_key,
                    DiseaseCase.location == region,
                )
                .first()
            )
            if existing:
                existing.case_count = cases_int
                existing.disease_name = disease_name_std
                existing.disease_type = "respiratory"
                existing.data_source = file.filename
                if note:
                    existing.note = note
                case_obj = existing
                updated += 1
            else:
                case_obj = DiseaseCase(
                    recorded_at=recorded_at,
                    icd_code=disease_key,
                    disease_name=disease_name_std,
                    disease_type="respiratory",
                    case_count=cases_int,
                    location=region,
                    data_source=file.filename,
                    note=note,
                    created_by=current_user.username,
                )
                db.add(case_obj)
                db.flush()
                imported += 1

            new_regions_to_register.add(region)

    else:
        # CSV bệnh viện: đếm SoTiepNhan phân biệt theo (ngày, mã ICD, khu vực).
        aggregate: dict[tuple, set[str]] = {}
        for row in reader:
            admit = (row.get("AdmissionDate") or "").strip()
            nhom = (row.get("NhomBenh") or "").strip()
            so_tiep_nhan = (row.get("SoTiepNhan") or "").strip()
            location = (row.get("District") or "Thành phố Hồ Chí Minh").strip()

            if not admit or not nhom or not so_tiep_nhan:
                skipped += 1
                continue

            disease_type = NHOM_BENH_MAPPING.get(nhom)
            if not disease_type:
                skipped += 1
                continue

            parsed = None
            for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%d/%m/%Y"):
                try:
                    parsed = datetime.strptime(admit, fmt)
                    break
                except ValueError:
                    continue
            if parsed is None:
                skipped += 1
                continue

            key = (parsed.date(), disease_type, location)
            aggregate.setdefault(key, set()).add(so_tiep_nhan)
            new_regions_to_register.add(location)

        for (rec_date, disease_type, location), ids in aggregate.items():
            recorded_at = datetime(rec_date.year, rec_date.month, rec_date.day)
            disease_name_std = ICD_TO_NAME.get(disease_type, disease_type)
            existing = (
                db.query(DiseaseCase)
                .filter(
                    DiseaseCase.recorded_at == recorded_at,
                    DiseaseCase.icd_code == disease_type,
                    DiseaseCase.location == location,
                )
                .first()
            )
            if existing:
                existing.case_count = len(ids)
                existing.disease_name = disease_name_std
                existing.data_source = file.filename
                updated += 1
            else:
                db.add(
                    DiseaseCase(
                        recorded_at=recorded_at,
                        icd_code=disease_type,
                        disease_name=disease_name_std,
                        disease_type="respiratory",
                        case_count=len(ids),
                        location=location,
                        data_source=file.filename,
                    )
                )
                imported += 1

    try:
        _register_pending_regions()
    except Exception as exc:
        logger.warning("Không thể tự đăng ký khu vực mới: %s", exc)

    try:
        db.commit()
    except Exception as exc:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to save: {exc}",
        )

    logger.info(
        f"Disease cases import by user={current_user.username}: "
        f"imported={imported} updated={updated} skipped={skipped}"
    )

    # Luôn None: DSS không suy luận độ nặng.
    auto_severity: dict | None = None

    return {
        "status": "ok",
        "imported": imported,
        "updated": updated,
        "skipped": skipped,
        "errors": errors[:200],
        "errors_truncated": len(errors) > 200,
        "auto_severity_recompute": auto_severity,
    }


@router.put("/{case_id}", response_model=DiseaseCaseResponse)
def update_disease_case(
    case_id: int,
    data: DiseaseCaseUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> Any:
    """Cập nhật một bản ghi số ca; chỉ ghi các trường được gửi, 409 nếu trùng tháng/mã/khu vực."""
    from app.models.disease_case import DiseaseCase
    from sqlalchemy import extract

    case = db.query(DiseaseCase).filter(DiseaseCase.id == case_id).first()
    if not case:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Disease case {case_id} not found",
        )

    new_recorded_at = data.recorded_at or case.recorded_at
    new_icd = data.icd_code or case.icd_code
    new_location = data.location or case.location

    duplicate = (
        db.query(DiseaseCase)
        .filter(
            DiseaseCase.id != case_id,
            extract("year", DiseaseCase.recorded_at) == new_recorded_at.year,
            extract("month", DiseaseCase.recorded_at) == new_recorded_at.month,
            DiseaseCase.icd_code == new_icd,
            DiseaseCase.location == new_location,
        )
        .first()
    )
    if duplicate:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail={
                "message": "Đã có bản ghi khác cho cùng tháng/bệnh/khu vực này.",
                "existing_id": duplicate.id,
            },
        )

    update_dict = data.model_dump(exclude_unset=True)
    for k, v in update_dict.items():
        setattr(case, k, v)

    try:
        db.commit()
        db.refresh(case)
    except Exception as exc:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to update: {exc}",
        )
    logger.info(f"Disease case {case_id} updated by user={current_user.username}")
    return case


@router.delete("/{case_id}", status_code=status.HTTP_200_OK)
def delete_disease_case(
    case_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> Any:
    """Xoá một bản ghi số ca."""
    from app.models.disease_case import DiseaseCase

    case = db.query(DiseaseCase).filter(DiseaseCase.id == case_id).first()
    if not case:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Disease case {case_id} not found",
        )

    db.delete(case)
    db.commit()
    logger.info(f"Disease case {case_id} deleted by user={current_user.username}")
    return {"message": "Disease case deleted", "id": case_id}


@router.get("/{case_id}/supply-usage")
def get_case_supply_usage(
    case_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> Any:
    """Nhu cầu thuốc ước tính cho một bản ghi số ca.
    Số lượng = case_count × norm_hieu_dung(i, g) (định mức thực nghiệm dss_demand); nhóm chưa có định mức → rỗng."""
    from app.models.disease_case import DiseaseCase
    from app.models.medical_supply import MedicalSupply
    from app.utils.icd_groups import nhom_cua_ma
    from app.services import dss_demand
    from sqlalchemy import extract

    case = db.query(DiseaseCase).filter(DiseaseCase.id == case_id).first()
    if not case:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Disease case {case_id} not found",
        )

    supplies: list[dict] = []
    severity_breakdown: dict | None = None

    khoi = nhom_cua_ma(case.icd_code)
    if khoi:
        try:
            emp_payload = dss_demand.norms_payload(db, khoi, limit=100000)
            emp_rows = emp_payload.get("rows", [])
        except Exception:
            emp_rows = []

        supply_codes = [r["supply_code"] for r in emp_rows if float(r.get("norm_hieu_dung") or 0) > 0]
        supply_by_code: dict[str, MedicalSupply] = {}
        if supply_codes:
            for s in (
                db.query(MedicalSupply)
                .filter(MedicalSupply.supply_code.in_(supply_codes))
                .all()
            ):
                supply_by_code[s.supply_code] = s

        for r in emp_rows:
            norm_hieu_dung = float(r.get("norm_hieu_dung") or 0)
            if norm_hieu_dung <= 0:
                continue
            code = r["supply_code"]
            supply = supply_by_code.get(code)
            used_qty = int(round(case.case_count * norm_hieu_dung))
            if used_qty <= 0:
                continue
            supplies.append({
                "supply_id": supply.id if supply else None,
                "code": code,
                "drug_code": supply.drug_code if supply else None,
                "name": (supply.ten_hoat_chat if supply else None) or r.get("ten") or code,
                "category": (
                    (supply.group_name or supply.category) if supply else None
                ) or r.get("nhom") or "Khác",
                "unit": (supply.unit if supply else None) or r.get("don_vi"),
                "description": supply.description if supply else None,
                "ratio": round(norm_hieu_dung, 4),
                "used_quantity": used_qty,
                "disease_label": case.disease_name,
                "khoi": khoi,
                "source": "norm",
            })

    supplies.sort(key=lambda s: s["used_quantity"], reverse=True)

    # So sánh với tháng trước cùng mã ICD, cùng khu vực.
    prev_year = case.recorded_at.year
    prev_month = case.recorded_at.month - 1
    if prev_month == 0:
        prev_month = 12
        prev_year -= 1
    prev_q = (
        db.query(func.coalesce(func.sum(DiseaseCase.case_count), 0))
        .filter(
            DiseaseCase.icd_code == case.icd_code,
            DiseaseCase.location == case.location,
            extract("year", DiseaseCase.recorded_at) == prev_year,
            extract("month", DiseaseCase.recorded_at) == prev_month,
        )
    )
    prev_total = int(prev_q.scalar() or 0)
    prev_supply_types = len(supplies)
    delta_cases = (
        round(((case.case_count - prev_total) / prev_total) * 100, 1)
        if prev_total > 0
        else None
    )

    return {
        "case": {
            "id": case.id,
            "recorded_at": case.recorded_at.isoformat(),
            "month": case.recorded_at.month,
            "year": case.recorded_at.year,
            "icd_code": case.icd_code,
            "disease_name": case.disease_name,
            # Trùng icd_code, giữ để tương thích schema.
            "disease_type": case.icd_code,
            "location": case.location,
            "case_count": case.case_count,
            "note": case.note,
        },
        "severity_breakdown": severity_breakdown,
        "summary": {
            "total_supply_types": len(supplies),
            "total_cases": case.case_count,
            "prev_total_cases": prev_total,
            "delta_cases_pct": delta_cases,
            "prev_total_supply_types": prev_supply_types,
        },
        "supplies": supplies,
    }


@router.get("/template")
async def download_template():
    """File CSV mẫu cho import số ca bệnh."""
    from fastapi.responses import StreamingResponse

    sample = (
        "month,disease_name,region,cases,note\n"
        "10/2024,Viêm phế quản cấp,TP. Hồ Chí Minh,145,\n"
        "10/2024,Nhiễm trùng đường hô hấp trên cấp,TP. Hồ Chí Minh,89,\n"
        "09/2024,Viêm họng cấp,Hà Nội,67,\n"
        "09/2024,Viêm xoang cấp,Hà Nội,42,Ghi chú tuỳ chọn\n"
    )
    return StreamingResponse(
        iter([sample]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=disease_cases_template.csv"},
    )


@router.get("/distinct-values")
async def distinct_values(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    """Giá trị phân biệt (nhóm ICD, mã ICD, khu vực) cho các bộ lọc trên giao diện."""
    from app.models.disease_case import DiseaseCase

    diseases_rows = (
        db.query(DiseaseCase.icd_code, DiseaseCase.disease_name)
        .distinct()
        .order_by(DiseaseCase.icd_code)
        .all()
    )
    regions = (
        db.query(DiseaseCase.location)
        .distinct()
        .order_by(DiseaseCase.location)
        .all()
    )

    icd_codes = [d[0] for d in diseases_rows if d[0]]
    diseases_list = [
        {"icd_code": icd, "disease_name": name or ICD_TO_NAME.get(icd, icd)}
        for icd, name in diseases_rows
        if icd
    ]
    # Chỉ liệt kê nhóm ICD có dữ liệu.
    disease_groups = []
    for khoa, ten in NHOM_ICD.items():
        ma_con = sorted({icd for icd in icd_codes if nhom_cua_ma(icd) == khoa})
        if ma_con:
            disease_groups.append({"key": khoa, "name": ten, "icd_codes": ma_con})

    return {
        "disease_groups": disease_groups,
        "icd_codes": icd_codes,
        "diseases": diseases_list,
        # Bí danh của icd_codes, giữ để tương thích schema.
        "disease_types": icd_codes,
        "regions": [r[0] for r in regions if r[0]],
    }
