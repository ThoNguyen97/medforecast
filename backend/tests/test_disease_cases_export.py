"""GET /api/v1/disease-cases/export — xuất Excel theo bộ lọc UI."""
import io
from datetime import datetime

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.dependencies import get_current_user
from app.main import app
from app.models.disease_case import DiseaseCase
from app.models.user import User

engine = create_engine(
    "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
)
Session = sessionmaker(autocommit=False, autoflush=False, bind=engine)


def _db():
    db = Session()
    try:
        yield db
    finally:
        db.close()


def _user():
    return User(id=1, username="tester", email="t@t.vn", password_hash="x",
                full_name="Tester", role="Administrator", is_active=True)


@pytest.fixture
def client():
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    db = Session()
    rows = [
        # (recorded_at, icd, name, group, location, cases)
        ("2025-07-01", "J20", "Viêm phế quản cấp tính", "J20-J22", "Thành phố Hồ Chí Minh", 14),
        ("2025-07-01", "J18", "Viêm phổi, tác nhân không xác định", "J09-J18", "Đồng Tháp", 11),
        ("2025-07-01", "J18", "Viêm phổi, tác nhân không xác định", None, "TP. Hồ Chí Minh", 5),
        ("2025-06-01", "J18", "Viêm phổi, tác nhân không xác định", "J09-J18", "Đồng Nai", 6),
        ("2025-08-01", "J06", "Nhiễm trùng hô hấp trên", "J00-J06", "Hà Nội", 21),
    ]
    for d, icd, name, grp, loc, n in rows:
        db.add(DiseaseCase(
            recorded_at=datetime.fromisoformat(d), recorded_date=datetime.fromisoformat(d).date(),
            icd_code=icd, disease_name=name, disease_group=grp, disease_type="respiratory",
            location=loc, case_count=n, data_source="test",
        ))
    db.commit()
    db.close()
    app.dependency_overrides[get_db] = _db
    app.dependency_overrides[get_current_user] = _user
    yield TestClient(app)
    Base.metadata.drop_all(bind=engine)


def _doc(resp):
    import openpyxl
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert 'filename="du_lieu_benh_' in resp.headers["content-disposition"]
    return openpyxl.load_workbook(io.BytesIO(resp.content))


def test_export_all(client):
    wb = _doc(client.get("/api/v1/disease-cases/export"))
    assert wb.sheetnames == ["Chi tiết", "Tổng hợp", "Bộ lọc"]
    ws = wb["Chi tiết"]
    assert [c.value for c in ws[1]][:8] == [
        "STT", "Tháng/Năm", "Mã nhóm", "Nhóm bệnh", "Mã ICD", "Tên bệnh", "Tỉnh/Thành", "Số ca mắc"]
    body = list(ws.iter_rows(min_row=2, values_only=True))
    data, tong = body[:-1], body[-1]
    assert len(data) == 5
    assert tong[6] == "TỔNG" and tong[7] == 57
    # Bản ghi thiếu disease_group được suy từ mã ICD; tên tỉnh được chuẩn hoá
    r = next(x for x in data if x[7] == 5)
    assert r[2] == "J09-J18" and r[3] == "Cúm và viêm phổi" and r[6] == "TP. Hồ Chí Minh"
    assert data[0][1] == "08/2025"  # sắp xếp mới nhất trước


def test_export_filter_group_and_location_alias(client):
    # location gửi tên chuẩn nhưng DB lưu cả 2 biến thể → phải bắt được cả 2
    wb = _doc(client.get("/api/v1/disease-cases/export",
                         params={"location": "TP. Hồ Chí Minh"}))
    data = list(wb["Chi tiết"].iter_rows(min_row=2, values_only=True))[:-1]
    assert sorted(x[7] for x in data) == [5, 14]

    wb = _doc(client.get("/api/v1/disease-cases/export",
                         params={"disease_group": "J09-J18"}))
    data = list(wb["Chi tiết"].iter_rows(min_row=2, values_only=True))[:-1]
    assert sorted(x[7] for x in data) == [5, 6, 11]


def test_export_filter_month_range_and_summary(client):
    wb = _doc(client.get("/api/v1/disease-cases/export",
                         params={"start_month": "2025-07", "end_month": "2025-06"}))  # nhập ngược
    data = list(wb["Chi tiết"].iter_rows(min_row=2, values_only=True))[:-1]
    assert sorted(x[7] for x in data) == [5, 6, 11, 14]
    th = list(wb["Tổng hợp"].iter_rows(min_row=2, values_only=True))
    # 2 dòng TP.HCM tháng 07 khác nhóm → không gộp; 2 biến thể tên → cùng 1 tỉnh
    assert th[-1][3] == "TỔNG" and th[-1][4] == 36
    assert len(th) - 1 == 4
    bl = dict(wb["Bộ lọc"].iter_rows(min_row=1, values_only=True))
    assert bl["Từ tháng"] == "07/2025" and bl["Đến tháng"] == "06/2025"
    assert bl["Số bản ghi"] == 4 and bl["Tổng số ca"] == 36


def test_export_invalid_month(client):
    r = client.get("/api/v1/disease-cases/export", params={"start_month": "07/2025"})
    assert r.status_code == 400


def test_export_requires_auth(client):
    app.dependency_overrides.pop(get_current_user, None)
    r = client.get("/api/v1/disease-cases/export")
    assert r.status_code in (401, 403)
